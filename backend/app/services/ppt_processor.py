"""PowerPoint (PPTX) text extraction with slide-level metadata."""
from pptx import Presentation
from typing import List, Dict, Any
import io
import logging

logger = logging.getLogger(__name__)


def extract_text_from_pptx(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """
    Extract text from PPTX slide by slide.
    Returns list of dicts: {filename, slide_number, title, text}
    """
    slides = []
    try:
        prs = Presentation(io.BytesIO(file_bytes))
        for i, slide in enumerate(prs.slides):
            texts = []
            title = ""
            
            for shape in slide.shapes:
                if shape.has_text_frame:
                    shape_text = shape.text_frame.text.strip()
                    if shape_text:
                        # First text often is title
                        if not title and len(shape_text) < 120:
                            title = shape_text
                        texts.append(shape_text)
            
            full_text = "\n".join(texts)
            full_text = clean_text(full_text)
            
            if full_text.strip():
                slides.append({
                    "filename": filename,
                    "slide_number": i + 1,
                    "title": title or f"Slide {i+1}",
                    "text": full_text,
                    "source_type": "pptx"
                })
        
        logger.info(f"Extracted {len(slides)} slides from {filename}")
    except Exception as e:
        logger.error(f"Failed to process PPTX {filename}: {e}")
        raise ValueError(f"Could not process PowerPoint: {str(e)}")
    
    return slides


def clean_text(text: str) -> str:
    if not text:
        return ""
    import re
    lines = [line.strip() for line in text.splitlines()]
    text = "\n".join(line for line in lines if line)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
